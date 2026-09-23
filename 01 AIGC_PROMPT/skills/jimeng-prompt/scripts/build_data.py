#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""构建即梦提示词数据集：清洗 Excel 并导出为 skill 自带的 JSON。

用法:
    python build_data.py
    python build_data.py --jimeng "D:/x/即梦800个神级指令合集.xlsx" --camera "D:/x/运镜方式.xlsx"
    python build_data.py --no-filter      # 保留原始全部条目（关闭内容安全过滤）

清洗内容:
    - 去除 UTF-8 BOM 字符 \\ufeff（表格中位置随机、可能重复出现）
    - 去除首尾空格、全角空格，内部换行/连续空白压缩为单个空格
    - 跳过空行与空单元格
    - 同一维度内去重并保持原有顺序
    - 丢弃命中内容安全词表的条目（构建时打印明细）
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 openpyxl，请先执行: python -m pip install openpyxl")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_DATA_DIR = SKILL_DIR / "data"

WORKSPACE = Path(r"F:\360MoveData\Users\27925\Desktop\Test agent\03  AIGC_Work\01 AIGC_PROMPT")
DEFAULT_JIMENG = WORKSPACE / "即梦800个神级指令合集.xlsx"
DEFAULT_CAMERA = WORKSPACE / "运镜方式.xlsx"

# key 与 Excel 列名的对应关系；character 维度由用户输入的主体替代，不参与随机
DIMENSIONS = [
    ("style", "风格与艺术表现"),
    ("lens", "镜头与构图"),
    ("character", "人物与服饰"),
    ("scene", "场景与背景"),
    ("detail", "细节与服饰"),
    ("tone", "色调与光效"),
    ("culture", "文化与环境"),
]

BOM = "\ufeff"

# 内容安全过滤：命中任一关键词的条目会被丢弃，避免产出被图片平台审核拦截的提示词
DEFAULT_FILTER = ("妓院", "裸体", "情色", "色情", "强奸", "幼女")


def clean(value: object) -> str:
    """清洗单个单元格文本。"""
    if value is None:
        return ""
    text = str(value).replace(BOM, "")
    text = text.replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip(" \t\r\n，,、；;|")


def find_header(rows, marker: str) -> tuple:
    """定位包含 marker 的表头行，返回 (行号, 清洗后的整行)。"""
    for idx, row in enumerate(rows):
        cleaned = [clean(cell) for cell in row]
        if marker in cleaned:
            return idx, cleaned
    raise RuntimeError(f"未能在表格中定位表头（缺少列名 {marker!r}），请检查 Excel 结构")


def dedupe(items):
    return list(dict.fromkeys(items))


def load_jimeng(path: Path, blocklist=DEFAULT_FILTER):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    header_idx, header = find_header(rows, "风格与艺术表现")
    col_of = {}
    for key, label in DIMENSIONS:
        if label not in header:
            raise RuntimeError(f"表头缺少列 {label!r}")
        col_of[key] = header.index(label)

    pools = {key: [] for key in col_of}
    dropped = {}
    for row in rows[header_idx + 1:]:
        for key, col in col_of.items():
            if col >= len(row):
                continue
            value = clean(row[col])
            if not value:
                continue
            hit = next((term for term in blocklist if term in value), None)
            if hit:
                dropped.setdefault(hit, []).append(value)
                continue
            pools[key].append(value)

    dimensions = []
    for key, label in DIMENSIONS:
        items = dedupe(pools[key])
        entry = {"key": key, "label": label, "count": len(items), "items": items}
        if key == "character":
            entry["replaced_by_subject"] = True
        dimensions.append(entry)

    workbook.close()
    data = {
        "source": path.name,
        "column_order": [label for _, label in DIMENSIONS],
        "note": "character（人物与服饰）由用户输入的主体替代，不参与随机抽取。",
        "dimensions": dimensions,
    }
    return data, dropped


def load_camera(path: Path, blocklist=DEFAULT_FILTER):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    header_idx, header = find_header(rows, "中文名称")
    col = header.index("中文名称")

    items = []
    for row in rows[header_idx + 1:]:
        if col < len(row):
            value = clean(row[col])
            if not value:
                continue
            if any(term in value for term in blocklist):
                continue
            items.append(value)

    workbook.close()
    items = dedupe(items)
    return {"source": path.name, "column": "中文名称", "count": len(items), "items": items}


def main() -> int:
    parser = argparse.ArgumentParser(description="清洗 Excel 并导出即梦提示词 JSON 数据集")
    parser.add_argument("--jimeng", default=str(DEFAULT_JIMENG), help="即梦800个神级指令合集.xlsx 路径")
    parser.add_argument("--camera", default=str(DEFAULT_CAMERA), help="运镜方式.xlsx 路径")
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="JSON 输出目录")
    parser.add_argument("--no-filter", action="store_true", help="关闭内容安全词过滤（默认开启）")
    args = parser.parse_args()

    jimeng_path = Path(args.jimeng)
    camera_path = Path(args.camera)
    for path in (jimeng_path, camera_path):
        if not path.exists():
            print(f"[错误] 找不到文件: {path}", file=sys.stderr)
            return 1

    data_dir = Path(args.data_dir)
    data_dir.mkdir(parents=True, exist_ok=True)

    blocklist = () if args.no_filter else DEFAULT_FILTER
    jimeng, dropped = load_jimeng(jimeng_path, blocklist)
    camera = load_camera(camera_path, blocklist)

    (data_dir / "jimeng_prompts.json").write_text(
        json.dumps(jimeng, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (data_dir / "camera_moves.json").write_text(
        json.dumps(camera, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    if dropped:
        print("[过滤] 命中内容安全词表，已丢弃以下条目：")
        for term, values in dropped.items():
            for value in values:
                print(f"  - {value}  (命中: {term})")

    print(f"[完成] 输出目录: {data_dir}")
    for dim in jimeng["dimensions"]:
        flag = "  (由主体替代)" if dim.get("replaced_by_subject") else ""
        print(f"  - {dim['label']}: {dim['count']} 条{flag}")
    print(f"  - 运镜: {camera['count']} 条 (默认不启用)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
